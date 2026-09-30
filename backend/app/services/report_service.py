"""Deterministic, evidence-bound local engineering artifact generation."""

from __future__ import annotations

import hashlib
import io
import json
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.core.config import settings

SUPPORTED_FORMATS = {"DOCX", "PDF", "XLSX", "PPTX"}
CLASSIFICATION_RANK = {"PUBLIC": 0, "INTERNAL": 1, "RESTRICTED": 2, "CONFIDENTIAL": 3, "SECRET": 4}


def _classification_allowed(classification: str, clearance: list[str]) -> bool:
    requested = CLASSIFICATION_RANK.get(str(classification).upper(), 4)
    return any(CLASSIFICATION_RANK.get(str(tag).upper(), -1) >= requested for tag in clearance)


class ReportService:
    def __init__(self, storage_dir: str | None = None) -> None:
        root = Path(storage_dir or settings.ARTIFACT_STORAGE_DIR).resolve()
        root.mkdir(parents=True, exist_ok=True)
        self.root = root

    @staticmethod
    def classification_allowed(classification: str, clearance: list[str]) -> bool:
        return _classification_allowed(classification, clearance)

    def validate_spec(self, spec: dict[str, Any], clearance: list[str]) -> dict[str, Any]:
        report_format = str(spec.get("format") or "").upper()
        if report_format not in SUPPORTED_FORMATS:
            raise ValueError("format must be DOCX, PDF, XLSX, or PPTX.")
        evidence = list(spec.get("evidence") or spec.get("evidence_blocks") or [])
        for item in evidence:
            classification = str(item.get("classification") or item.get("classification_tag") or spec.get("classification") or "INTERNAL").upper()
            if not _classification_allowed(classification, clearance):
                raise PermissionError("Report contains evidence above the user's clearance.")
        classification = str(spec.get("classification") or "INTERNAL").upper()
        if not _classification_allowed(classification, clearance):
            raise PermissionError("Report classification exceeds the user's clearance.")
        spec = dict(spec)
        spec["format"] = report_format
        spec["classification"] = classification
        spec["evidence"] = evidence
        spec["review_status"] = "REVIEW_REQUIRED" if spec.get("recommended_verification") or spec.get("recommended_action") or spec.get("operational_recommendation") else "DRAFT"
        return spec

    def generate(self, spec: dict[str, Any], created_by: str, clearance: list[str], agent_run_id: str | None = None) -> dict[str, Any]:
        spec = self.validate_spec(spec, clearance)
        artifact_id = f"ART-{uuid.uuid4().hex}"
        extension = spec["format"].lower()
        storage_name = f"{artifact_id}.{extension}"
        content = self._render(spec)
        path = (self.root / storage_name).resolve()
        path.write_bytes(content)
        digest = hashlib.sha256(content).hexdigest()
        return {
            "artifact_id": artifact_id,
            "report_type": str(spec.get("report_type") or "ENGINEERING_REPORT"),
            "format": spec["format"],
            "storage_name": storage_name,
            "created_by": created_by,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "agent_run_id": agent_run_id,
            "classification": spec["classification"],
            "source_count": len(spec["evidence"]),
            "status": spec["review_status"],
            "content_hash": digest,
            "metadata": {"title": spec.get("title") or "Aegis Engineering Report"},
            "path": str(path),
            "spec": spec,
        }

    def _render(self, spec: dict[str, Any]) -> bytes:
        fmt = spec["format"]
        if fmt == "DOCX":
            return self._docx(spec)
        if fmt == "PDF":
            return self._pdf(spec)
        if fmt == "XLSX":
            return self._xlsx(spec)
        return self._pptx(spec)

    @staticmethod
    def _sections(spec: dict[str, Any]) -> list[tuple[str, Any]]:
        sections = [("Executive Summary", spec.get("executive_summary")), ("Asset / Equipment", spec.get("asset"))]
        sections.extend([("Findings", spec.get("findings")), ("Evidence", spec.get("evidence"))])
        sections.extend([("Possible Interpretation", spec.get("possible_interpretation")), ("Recommended Verification", spec.get("recommended_verification"))])
        sections.extend([("Calculations", spec.get("calculations")), ("Review / Approval Status", spec.get("review_status"))])
        return [(title, value) for title, value in sections if value not in (None, "", [], {})]

    @staticmethod
    def _evidence_lines(spec: dict[str, Any]) -> list[str]:
        lines = []
        for item in spec.get("evidence", []):
            source = item.get("document_name") or item.get("document") or item.get("document_id") or "Source"
            page = item.get("page") or item.get("location", {}).get("page") or "N/A"
            provenance = ", ".join(f"{key}={item[key]}" for key in ("document_id", "source_type", "extraction_method", "classification", "retrieval_reason") if item.get(key) is not None)
            lines.append(f"{source} | Page {page} | {provenance} | {item.get('snippet') or item.get('text') or ''}")
        return lines

    def _docx(self, spec: dict[str, Any]) -> bytes:
        from docx import Document
        document = Document()
        document.add_heading(str(spec.get("title") or "Aegis Engineering Report"), 0)
        document.add_paragraph(f"Classification: {spec['classification']} | Status: {spec['review_status']}")
        for title, value in self._sections(spec):
            document.add_heading(title, level=1)
            if title == "Evidence":
                table = document.add_table(rows=1, cols=2)
                table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Source / Page", "Provenance / Evidence"
                for line in self._evidence_lines(spec):
                    cells = table.add_row().cells
                    parts = line.split(" | ", 1)
                    cells[0].text, cells[1].text = parts[0], parts[1] if len(parts) > 1 else ""
            else:
                document.add_paragraph(json.dumps(value, indent=2, default=str) if isinstance(value, (dict, list)) else str(value))
        if spec["review_status"] == "REVIEW_REQUIRED":
            document.add_paragraph("HUMAN REVIEW REQUIRED")
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _pdf(self, spec: dict[str, Any]) -> bytes:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=letter)
        y = 760
        pdf.setTitle(str(spec.get("title") or "Aegis Engineering Report"))
        for title, value in [("Title", spec.get("title") or "Aegis Engineering Report"), ("Classification", spec["classification"]), *self._sections(spec)]:
            if y < 60:
                pdf.showPage(); y = 760
            pdf.setFont("Helvetica-Bold", 11); pdf.drawString(48, y, title); y -= 16
            pdf.setFont("Helvetica", 9)
            text = self._evidence_lines(spec) if title == "Evidence" else [json.dumps(value, default=str) if isinstance(value, (dict, list)) else str(value)]
            for line in text:
                for fragment in [line[i:i + 105] for i in range(0, len(line), 105)] or [""]:
                    if y < 45: pdf.showPage(); y = 760
                    pdf.drawString(56, y, fragment); y -= 12
            y -= 8
        if spec["review_status"] == "REVIEW_REQUIRED":
            pdf.setFont("Helvetica-Bold", 11); pdf.drawString(48, max(y, 40), "HUMAN REVIEW REQUIRED")
        pdf.save()
        return output.getvalue()

    def _xlsx(self, spec: dict[str, Any]) -> bytes:
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Engineering Evidence"
        headers = ["Asset", "Tag", "Measurement", "Unit", "Date", "Source", "Page", "Classification"]
        sheet.append(headers)
        for item in spec["evidence"]:
            tags = item.get("tag") or item.get("engineering_tags") or ""
            if isinstance(tags, list):
                tags = ", ".join(str(tag) for tag in tags)
            sheet.append([spec.get("asset", ""), tags, item.get("measurement", ""), item.get("unit", ""), item.get("date", ""), item.get("document_name") or item.get("document_id", ""), item.get("page") or item.get("location", {}).get("page", ""), item.get("classification", spec["classification"])])
        output = io.BytesIO(); workbook.save(output); return output.getvalue()

    def _pptx(self, spec: dict[str, Any]) -> bytes:
        from pptx import Presentation
        presentation = Presentation()
        for title, value in [("Problem / Asset", spec.get("asset")), *self._sections(spec)]:
            if value in (None, "", [], {}): continue
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            body = slide.placeholders[1].text_frame
            values = self._evidence_lines(spec) if title == "Evidence" else [str(value)]
            body.text = "\n".join(values)
        if spec["review_status"] == "REVIEW_REQUIRED":
            slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "Review Status"; slide.placeholders[1].text = "HUMAN REVIEW REQUIRED"
        output = io.BytesIO(); presentation.save(output); return output.getvalue()


report_service = ReportService()
